import io
import openpyxl
from openpyxl.styles import PatternFill
import pandas as pd
import requests
import streamlit as st

# Configuración de la página
st.set_page_config(
    page_title="Calculadora de Envíos", page_icon="💸", layout="centered"
)

st.title("💸 Calculadora de Depósitos en Bolívares")
st.write(
    "Obtén las tasas del día en tiempo real y genera tu archivo Excel de cálculo."
)


# --- 1. OBTENCIÓN DE TASAS ---
@st.cache_data(ttl=600)  # Guarda en caché por 10 minutos para ser veloz
def obtener_tasas():
    try:
        r_bcv = requests.get("https://ve.dolarapi.com/v1/dolares/oficial", timeout=5).json()
        bcv = r_bcv.get("promedio", 0.0)
    except Exception:
        bcv = 0.0

    try:
        r_par = requests.get("https://ve.dolarapi.com/v1/dolares/paralelo", timeout=5).json()
        paralelo = r_par.get("promedio", 0.0)
    except Exception:
        paralelo = 0.0

    try:
        r_bna = requests.get("https://dolarapi.com/v1/dolares/oficial", timeout=5).json()
        bna_venta = r_bna.get("venta", 0.0)
    except Exception:
        bna_venta = 0.0

    return bcv, paralelo, bna_venta


bcv, paralelo, bna_venta = obtener_tasas()

# Mostrar tarjetas con las tasas
st.subheader("📊 Tasas en tiempo real")
col1, col2, col3 = st.columns(3)
col1.metric("BCV", f"Bs. {bcv:,.2f}" if bcv else "N/D")
col2.metric("Paralelo", f"Bs. {paralelo:,.2f}" if paralelo else "N/D")
col3.metric("BNA Venta", f"$ {bna_venta:,.2f}" if bna_venta else "N/D")

st.divider()

# --- 2. FORMULARIO DE ENTRADA DE DATOS ---
st.subheader("📝 Ingresa los montos a calcular")

# Selección de la tasa USD/Bs a utilizar
tasa_usd_bs_default = paralelo if paralelo > 0 else bcv
tasa_usd_bs = st.number_input(
    "Tasa USD a Bs. a utilizar:",
    value=float(tasa_usd_bs_default),
    format="%.4f",
)

tasa_envios = st.number_input(
    "Tasa de Envíos JL Express (Bs./ARS):", value=0.5205, format="%.4f"
)

col_usd, col_bs, col_ars = st.columns(3)
with col_usd:
    monto_usd = st.number_input("Monto en USD:", value=200.0, step=10.0)
with col_bs:
    monto_bs = st.number_input("Monto en Bs:", value=17000.0, step=500.0)
with col_ars:
    monto_ars = st.number_input("Monto en ARS:", value=0.0, step=1000.0)


# --- 3. PROCESAMIENTO Y GENERACIÓN DEL EXCEL ---
def generar_excel(
    tasa_usd_bs, tasa_envios, monto_usd, monto_bs, monto_ars, plantilla_path
):
    wb = openpyxl.load_workbook(plantilla_path)
    sheet = wb["Hoja1"]

    # 1. Asegurar valores numéricos limpios
    t_usd_bs = float(tasa_usd_bs) if tasa_usd_bs else 0.0
    t_envios = float(tasa_envios) if tasa_envios else 0.0

    m_usd = float(monto_usd) if (monto_usd and float(monto_usd) > 0) else 0.0
    m_bs = float(monto_bs) if (monto_bs and float(monto_bs) > 0) else 0.0
    m_ars = float(monto_ars) if (monto_ars and float(monto_ars) > 0) else 0.0

    usd_ars = (t_usd_bs / t_envios) if t_envios > 0 else 0.0

    # 2. Celdas de origen (Tabla izquierda)
    sheet["B4"] = t_usd_bs
    sheet["B3"] = t_envios
    sheet["C6"] = m_usd
    sheet["C7"] = m_bs
    sheet["C8"] = m_ars

    # 3. Cálculos de conversión para la Tabla de Resultado (Derecha)
    # Fila 4 (USD)
    f4_bs = m_usd * t_usd_bs
    g4_ars = m_usd * usd_ars
    h4_usd = m_usd

    # Fila 5 (Bs)
    f5_bs = m_bs
    g5_ars = (m_bs / t_envios) if t_envios > 0 else 0.0
    h5_usd = (m_bs / t_usd_bs) if t_usd_bs > 0 else 0.0

    # Fila 6 (ARS)
    f6_bs = m_ars * t_envios
    g6_ars = m_ars
    h6_usd = (m_ars / usd_ars) if usd_ars > 0 else 0.0

    # Escribir los valores numéricos
    sheet["F4"], sheet["G4"], sheet["H4"] = f4_bs, g4_ars, h4_usd
    sheet["F5"], sheet["G5"], sheet["H5"] = f5_bs, g5_ars, h5_usd
    sheet["F6"], sheet["G6"], sheet["H6"] = f6_bs, g6_ars, h6_usd

    # Totales (Fila 7)
    sheet["F7"] = f4_bs + f5_bs + f6_bs
    sheet["G7"] = g4_ars + g5_ars + g6_ars
    sheet["H7"] = h4_usd + h5_usd + h6_usd

    # 4. Estilos de Relleno (Resaltado dinámico según combinaciones)
    fill_destacado = PatternFill(
        start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
    )
    fill_limpio = PatternFill(fill_type=None)  # Fondo transparente / blanco

    # H4 se pinta ÚNICAMENTE si entró dinero en USD
    sheet["H4"].fill = fill_destacado if m_usd > 0 else fill_limpio

    # F5 se pinta ÚNICAMENTE si entró dinero en Bolívares
    sheet["F5"].fill = fill_destacado if m_bs > 0 else fill_limpio

    # G6 se pinta ÚNICAMENTE si entró dinero en Pesos ARS
    sheet["G6"].fill = fill_destacado if m_ars > 0 else fill_limpio

    # Guardar en memoria
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

plantilla_path = "Cálculo rápido para depósitos en bolívares.xlsx"

if st.button("🚀 Calcular y Preparar Excel", type="primary"):
    try:
        excel_buffer = generar_excel(
            tasa_usd_bs,
            tasa_envios,
            monto_usd,
            monto_bs,
            monto_ars,
            plantilla_path,
        )

        st.success("¡Cálculo realizado con éxito!")

        # Botón para descargar el archivo Excel generado
        st.download_button(
            label="📥 Descargar Excel Calculado",
            data=excel_buffer,
            file_name="Calculo_Depositos_Actualizado.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as e:
        st.error(f"Error al procesar el Excel: {e}")